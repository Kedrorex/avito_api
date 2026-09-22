import openpyxl
import sys
import pathlib
import json

# Find the actual file
fid_dir = pathlib.Path('fid')
output = []

for f in fid_dir.iterdir():
    output.append(f"Found: {f.name}")
    wb = openpyxl.load_workbook(str(f))
    ws = wb.active
    output.append(f'Sheet: {ws.title}')
    output.append(f'Max row: {ws.max_row}')
    output.append(f'Max col: {ws.max_column}')
    output.append('')
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        output.append(json.dumps(vals, ensure_ascii=False))

# Write to file with UTF-8
with open('template_output.txt', 'w', encoding='utf-8') as fout:
    fout.write('\n'.join(output))

print("Output written to template_output.txt")
